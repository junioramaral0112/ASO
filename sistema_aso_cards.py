import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
import os
from PIL import Image
import base64
import unicodedata
from urllib.parse import unquote

# =========================================================
# CONFIGURAÇÃO DE CAMINHOS DINÂMICOS
# =========================================================

BASE_DIR = os.path.dirname(__file__) if "__file__" in locals() else "."

def localizar_arquivo(caminho_local, nome_arquivo):
    if os.path.exists(caminho_local):
        return caminho_local
    caminho_projeto = os.path.join(BASE_DIR, nome_arquivo)
    return caminho_projeto if os.path.exists(caminho_projeto) else None

# Caminhos dos arquivos
PATH_ASO_IMG = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\ASO.png", "ASO.png")
PATH_LOGO = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\logo.png", "logo.png")
PATH_LOGO_DOC = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\adivitta.png", "adivitta.png")

# Função para converter imagem para Base64
def carregar_imagem_base64(path):
    if path and os.path.exists(path):
        with open(path, "rb") as img_file:
            return base64.b64encode(img_file.read()).decode()
    return None

def normalizar_texto(texto):
    if not isinstance(texto, str):
        return ""
    texto = unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("ASCII")
    return texto.strip().upper()

# =========================================================
# CONFIGURAÇÃO DA PÁGINA
# =========================================================

try:
    if PATH_ASO_IMG:
        page_icon = Image.open(PATH_ASO_IMG)
    else:
        page_icon = "🛡️"
except:
    page_icon = "🛡️"

st.set_page_config(
    page_title="Gestão ASO Pro",
    layout="wide",
    page_icon=page_icon
)

# =========================================================
# CSS GLOBAL E TOPO PERSONALIZADO
# =========================================================

aso_base64 = carregar_imagem_base64(PATH_ASO_IMG)

st.markdown(f"""
<style>
/* ESCONDER NAVEGAÇÃO AUTOMÁTICA */
[data-testid="stSidebarNav"] {{display: none !important;}}

/* FUNDO E ESTILOS GERAIS */
.stApp {{ background-color: #f1f5f9; }}
section[data-testid="stSidebar"] {{ background: linear-gradient(180deg, #8390a8, #1e293b); }}
section[data-testid="stSidebar"] label, section[data-testid="stSidebar"] p {{ color: white !important; }}
div[data-testid="metric-container"] {{ background: white; border-radius: 18px; padding: 15px; border: 1px solid #e2e8f0; box-shadow: 0 4px 18px rgba(0,0,0,0.06); }}

/* BOTÃO DE DOWNLOAD (PRINCIPAL) */
.stDownloadButton button {{ width: 100%; background: linear-gradient(90deg, #2563eb, #1d4ed8); color: white; border-radius: 12px; font-weight: bold; }}

/* AJUSTE DO BOTÃO CHECKLIST NA SIDEBAR (RESOLVE TEXTO APAGADO) */
section[data-testid="stSidebar"] .stButton > button {{
    background-color: #2563eb !important;
    color: white !important;
    border-radius: 12px !important;
    border: 1px solid rgba(255,255,255,0.3) !important;
    font-weight: bold !important;
    height: 45px !important;
    transition: 0.3s;
}}

section[data-testid="stSidebar"] .stButton > button:hover {{
    background-color: #1d4ed8 !important;
    border: 1px solid #f9cc0b !important;
}}

.footer {{ position: fixed; left: 0; bottom: 0; width: 100%; background: rgba(0,0,0,0.8); color: white; text-align: center; padding: 10px; font-size: 13px; z-index: 999; }}

.menu-titulo {{
    color: white;
    font-weight: bold;
    font-size: 14px;
    margin-top: 20px;
    margin-bottom: 10px;
    border-bottom: 1px solid rgba(255,255,255,0.2);
    padding-bottom: 5px;
}}

/* TÍTULO COM IMAGEM AMPLIADA */
.header-wrapper {{
    display: flex;
    align-items: center;
    gap: 20px;
    margin-bottom: 5px;
}}
.main-title {{
    font-size: 40px;
    font-weight: 800;
    color: #111827;
    margin: 0;
}}
.sub-title {{
    color: #64748b;
    margin-bottom: 25px;
    margin-left: 120px;
}}
</style>

<div class="header-wrapper">
    {"<img src='data:image/png;base64," + aso_base64 + "' width='100'>" if aso_base64 else "🛡️"}
    <h1 class="main-title">Gestão ASO Pro</h1>
</div>
<div class="sub-title">Sistema Inteligente de Gestão de ASO</div>
""", unsafe_allow_html=True)

# =========================================================
# LÓGICA DE DADOS
# =========================================================

SHEET_ID = "1G_oVT9gK-n_jGh5R4g65qUwK_MfQGvCX-SA4NHNNflU"
UNIDADES = {
    "D-ITU": "1323067532", "D-MG": "1071212860", 
    "J-CHP": "1549718037", "S-SJ": "1712391604", "J-CTBA": "145843404"
}

@st.cache_data(ttl=300)
def load_data(gid):
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=csv&gid={gid}"
    df = pd.read_csv(url)
    df.columns = df.columns.str.strip()
    return df

def gerar_docx(dados, tipo, data_sugestao):
    doc = Document()
    if PATH_LOGO_DOC and os.path.exists(PATH_LOGO_DOC):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(PATH_LOGO_DOC, width=Inches(1.2))

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run(f"FORMULÁRIO PARA AGENDAMENTO {tipo}")
    run.bold, run.font.size = True, Pt(16)

    table = doc.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    labels = ["Nome Completo", "Cargo", "Setor", "Unidade", "Cliente", "Local", "Data Sugestão"]
    valores = [str(dados["Nome"]), str(dados["Cargo"]), str(dados["Setor"]), str(dados.get("UNIDADE", aba_nome)), "MACROMAQ", "Arapoti", data_sugestao.strftime("%d/%m/%Y")]

    for i, (l, v) in enumerate(zip(labels, valores)):
        table.rows[i].cells[0].text, table.rows[i].cells[1].text = l, v
        table.rows[i].cells[0].paragraphs[0].runs[0].bold = True

    target = BytesIO()
    doc.save(target)
    return target.getvalue()

# =========================================================
# CAPTURA DE PARÂMETROS DA URL (?colaborador=...)
# =========================================================
colab_param = st.query_params.get("colaborador", None)
unidade_inicial_idx = 0

# Se veio um colaborador pela URL, procura em qual unidade ele está
if colab_param:
    nome_alvo = normalizar_texto(unquote(str(colab_param)))
    for idx_u, (nome_unidade, gid_unidade) in enumerate(UNIDADES.items()):
        try:
            df_temp = load_data(gid_unidade)
            if not df_temp.empty and "Nome" in df_temp.columns:
                nomes_norm = df_temp["Nome"].astype(str).apply(normalizar_texto)
                if any(nome_alvo in n or n in nome_alvo for n in nomes_norm):
                    unidade_inicial_idx = idx_u
                    break
        except:
            pass

# =========================================================
# SIDEBAR
# =========================================================

if PATH_LOGO and os.path.exists(PATH_LOGO):
    st.sidebar.image(PATH_LOGO, width=300)

aba_nome = st.sidebar.selectbox(
    "🏢 Selecione a unidade",
    list(UNIDADES.keys()),
    index=unidade_inicial_idx
)

st.sidebar.success("✅ Sistema Online")

st.sidebar.markdown("<br>", unsafe_allow_html=True)
st.sidebar.markdown(
    "<div class='menu-titulo'>⚡ ACESSO RÁPIDO</div>",
    unsafe_allow_html=True
)

if st.sidebar.button("📋 CHECKLIST SSMA"):
    st.switch_page("pages/app_ssma_ia.py")

# =========================================================
# PROCESSAMENTO PRINCIPAL
# =========================================================
try:
    df = load_data(UNIDADES[aba_nome])
    if not df.empty:
        hoje = datetime.now()
        df["Venc"] = pd.to_datetime(df["Venc"], dayfirst=True, errors="coerce")
        df = df.dropna(subset=["Venc"])
        df["Dias"] = (df["Venc"] - hoje).dt.days
        
        # Filtro prioritário caso tenha vindo pela URL
        registro_focado = None
        if colab_param:
            nome_alvo = normalizar_texto(unquote(str(colab_param)))
            df["Nome_Norm"] = df["Nome"].astype(str).apply(normalizar_texto)
            match_colab = df[df["Nome_Norm"].str.contains(nome_alvo, na=False) | df["Nome_Norm"].apply(lambda x: x in nome_alvo)]
            if not match_colab.empty:
                registro_focado = match_colab.iloc[0]

        alertas = df[df["Venc"] <= hoje + timedelta(days=10)].copy().sort_values(by="Venc")

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("🏢 Unidade", aba_nome); c2.metric("👥 Colaboradores", len(df))
        c3.metric("⚠️ Pendentes", len(alertas)); c4.metric("🚨 Vencidos", len(alertas[alertas["Dias"] < 0]))

        # SE HOUVER UM COLABORADOR SELECIONADO PELA URL, MOSTRA ELE EM DESTAQUE NO TOPO
        if registro_focado is not None:
            st.markdown("---")
            st.markdown(f"### 🎯 Colaborador Selecionado: **{registro_focado['Nome']}**")
            
            dias_foco = int(registro_focado["Dias"])
            cor_foco, fundo_foco, status_foco = ("#ef4444", "#fff1f2", "🚨 ASO VENCIDO") if dias_foco < 0 else (("#f59e0b", "#fff7ed", f"⚠️ Vence em {dias_foco} dias") if dias_foco <= 3 else ("#10b981", "#ecfdf5", f"✅ Vence em {dias_foco} dias"))
            
            col_card, col_form = st.columns([1, 1])
            with col_card:
                html_card_foco = f"""
                <div style="background:{fundo_foco}; border-left:8px solid {cor_foco}; border-radius:18px; padding:22px; margin-bottom:15px; box-shadow:0 4px 18px rgba(0,0,0,0.08); font-family:Arial;">
                    <div style="font-size:22px; font-weight:700; color:#111827; margin-bottom:10px;">{registro_focado['Nome']}</div>
                    <div style="color:#475569; font-size:15px; line-height:1.8;">👔 <b>Cargo:</b> {registro_focado['Cargo']}<br>🏭 <b>Setor:</b> {registro_focado['Setor']}<br>📅 <b>Vencimento:</b> {registro_focado['Venc'].strftime('%d/%m/%Y')}</div>
                    <div style="margin-top:15px; font-size:16px; font-weight:bold; color:{cor_foco};">{status_foco}</div>
                </div>"""
                components.html(html_card_foco, height=220)
                
            with col_form:
                with st.container():
                    st.markdown("##### 📄 Formulário de Agendamento")
                    tipo_foco = st.selectbox("Tipo de Exame", ["PERIÓDICO", "MUDANÇA DE RISCO", "RETORNO"], key="tipo_foco")
                    dt_foco = st.date_input("Data sugerida", value=hoje + timedelta(days=2), key="data_foco")
                    btn_doc_foco = gerar_docx(registro_focado, tipo_foco, dt_foco)
                    st.download_button(label="📥 Baixar Documento ASO", data=btn_doc_foco, file_name=f"ASO_{registro_focado['Nome']}.docx", key="btn_foco")

        st.markdown("---")
        st.markdown("#### 📋 Alertas de Vencimento da Unidade")
        cols = st.columns(2)
        for idx, (_, row) in enumerate(alertas.iterrows()):
            col = cols[idx % 2]
            dias = int(row["Dias"])
            cor, fundo, status = ("#ef4444", "#fff1f2", "🚨 ASO VENCIDO") if dias < 0 else (("#f59e0b", "#fff7ed", f"⚠️ Vence em {dias} dias") if dias <= 3 else ("#10b981", "#ecfdf5", f"✅ Vence em {dias} dias"))

            html_card = f"""
            <div style="background:{fundo}; border-left:8px solid {cor}; border-radius:18px; padding:22px; margin-bottom:15px; box-shadow:0 4px 18px rgba(0,0,0,0.08); font-family:Arial;">
                <div style="font-size:22px; font-weight:700; color:#111827; margin-bottom:10px;">{row['Nome']}</div>
                <div style="color:#475569; font-size:15px; line-height:1.8;">👔 <b>Cargo:</b> {row['Cargo']}<br>🏭 <b>Setor:</b> {row['Setor']}<br>📅 <b>Vencimento:</b> {row['Venc'].strftime('%d/%m/%Y')}</div>
                <div style="margin-top:15px; font-size:16px; font-weight:bold; color:{cor};">{status}</div>
            </div>"""

            with col:
                components.html(html_card, height=250)
                with st.expander(f"📄 Gerar Agendamento - {row['Nome'].split()[0]}"):
                    tipo = st.selectbox("Tipo de Exame", ["PERIÓDICO", "MUDANÇA DE RISCO", "RETORNO"], key=f"t_{idx}")
                    dt_s = st.date_input("Data sugerida", value=hoje + timedelta(days=2), key=f"d_{idx}")
                    btn_doc = gerar_docx(row, tipo, dt_s)
                    st.download_button(label="📥 Baixar Documento", data=btn_doc, file_name=f"ASO_{row['Nome']}.docx", key=f"b_{idx}")
except Exception as e:
    st.error(f"Erro: {e}")

st.markdown("""<div class="footer">© 2026 Gestão Documentos | Desenvolvido por: Dilceu Junior</div>""", unsafe_allow_html=True)
