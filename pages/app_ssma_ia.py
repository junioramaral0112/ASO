import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
import os
import json
import urllib.request
from PIL import Image
import base64
import unicodedata
from urllib.parse import unquote

# =========================================================
# CONFIGURAÇÃO DE CAMINHOS E WEBHOOK DO APPS SCRIPT
# =========================================================

BASE_DIR = os.path.dirname(__file__) if "__file__" in locals() else "."

APPS_SCRIPT_WEBHOOK_URL = "https://script.google.com/a/macros/macromaq.com.br/s/AKfycbyNCchIjNTiUna4cWlyDOJr8KYqpYDBdNqormPb91d_xYnOu7L1lxrZDAL4q0I2lwux/exec"

def localizar_arquivo(caminho_local, nome_arquivo):
    if os.path.exists(caminho_local):
        return caminho_local
    caminho_projeto = os.path.join(BASE_DIR, nome_arquivo)
    return caminho_projeto if os.path.exists(caminho_projeto) else None

PATH_ASO_IMG = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\ASO.png", "ASO.png")
PATH_LOGO = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\logo.png", "logo.png")
PATH_LOGO_DOC = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\adivitta.png", "adivitta.png")
PATH_AGENDAMENTOS_JSON = os.path.join(BASE_DIR, "agendamentos_persistentes.json")

def carregar_imagem_base64(path):
    if path and os.path.exists(path):
        with open(path, "rb") as img_file:
            return base64.b64encode(img_file.read()).decode()
    return None

def normalizar_texto(texto):
    if not isinstance(texto, str):
        return ""
    texto = unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("ASCII")
    return texto.strip().upper()

# =========================================================
# PERSISTÊNCIA LOCAL E ENVIO DE DADOS (APPS SCRIPT)
# =========================================================

def ler_banco_agendamentos():
    if os.path.exists(PATH_AGENDAMENTOS_JSON):
        try:
            with open(PATH_AGENDAMENTOS_JSON, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def gravar_agendamento(chave, dados):
    banco = ler_banco_agendamentos()
    banco[chave] = dados
    try:
        with open(PATH_AGENDAMENTOS_JSON, "w", encoding="utf-8") as f:
            json.dump(banco, f, ensure_ascii=False, indent=2)
    except Exception as e:
        st.error(f"Erro ao salvar localmente: {e}")

def enviar_webhook_planilha(payload):
    """Envia requisição para o Google Apps Script e trata redirecionamento."""
    if not APPS_SCRIPT_WEBHOOK_URL or not APPS_SCRIPT_WEBHOOK_URL.startswith("http"):
        return {"status": "erro", "mensagem": "URL do Apps Script não preenchida em APPS_SCRIPT_WEBHOOK_URL"}
    try:
        dados_json = json.dumps(payload).encode("utf-8")
        req = urllib.request.Request(
            APPS_SCRIPT_WEBHOOK_URL,
            data=dados_json,
            headers={"Content-Type": "application/json"}
        )
        opener = urllib.request.build_opener()
        with opener.open(req, timeout=12) as response:
            res_text = response.read().decode("utf-8")
            try:
                return json.loads(res_text)
            except Exception:
                return {"status": "sucesso"}
    except Exception as ex:
        return {"status": "erro", "mensagem": str(ex)}

def remover_agendamento(chave, payload_sync=None):
    banco = ler_banco_agendamentos()
    banco[chave] = {"status": False, "data": "", "hora": "", "obs": ""}
    try:
        with open(PATH_AGENDAMENTOS_JSON, "w", encoding="utf-8") as f:
            json.dump(banco, f, ensure_ascii=False, indent=2)
    except Exception as e:
        st.error(f"Erro ao atualizar: {e}")
    if payload_sync:
        enviar_webhook_planilha(payload_sync)

# =========================================================
# CONFIGURAÇÃO DA PÁGINA
# =========================================================

try:
    if PATH_ASO_IMG:
        page_icon = Image.open(PATH_ASO_IMG)
    else:
        page_icon = "🛡️"
except Exception:
    page_icon = "🛡️"

st.set_page_config(
    page_title="Gestão ASO Pro",
    layout="wide",
    page_icon=page_icon
)

# =========================================================
# CSS GLOBAL E CABEÇALHO
# =========================================================

aso_base64 = carregar_imagem_base64(PATH_ASO_IMG)

st.markdown(f"""
<style>
[data-testid="stSidebarNav"] {{display: none !important;}}
.stApp {{ background-color: #f1f5f9; }}
section[data-testid="stSidebar"] {{ background: linear-gradient(180deg, #8390a8, #1e293b); }}
section[data-testid="stSidebar"] label, section[data-testid="stSidebar"] p {{ color: white !important; }}
div[data-testid="metric-container"] {{ background: white; border-radius: 18px; padding: 15px; border: 1px solid #e2e8f0; box-shadow: 0 4px 18px rgba(0,0,0,0.06); }}

.stDownloadButton button {{ width: 100%; background: linear-gradient(90deg, #2563eb, #1d4ed8); color: white; border-radius: 12px; font-weight: bold; }}

section[data-testid="stSidebar"] .stButton > button {{
    background-color: #2563eb !important;
    color: white !important;
    border-radius: 12px !important;
    border: 1px solid rgba(255,255,255,0.3) !important;
    font-weight: bold !important;
    height: 45px !important;
    transition: 0.3s;
}}

.footer {{ position: fixed; left: 0; bottom: 0; width: 100%; background: rgba(0,0,0,0.8); color: white; text-align: center; padding: 10px; font-size: 13px; z-index: 999; }}

.menu-titulo {{
    color: white;
    font-weight: bold;
    font-size: 14px;
    margin-top: 20px;
    margin-bottom: 10px;
    border-bottom: 1px solid rgba(255,255,255,0.2);
    padding-bottom: 5px;
}}

.header-wrapper {{
    display: flex;
    align-items: center;
    gap: 20px;
    margin-bottom: 5px;
}}
.main-title {{
    font-size: 40px;
    font-weight: 800;
    color: #111827;
    margin: 0;
}}
.sub-title {{
    color: #64748b;
    margin-bottom: 25px;
    margin-left: 120px;
}}
</style>

<div class="header-wrapper">
    {"<img src='data:image/png;base64," + aso_base64 + "' width='100'>" if aso_base64 else "🛡️"}
    <h1 class="main-title">Gestão ASO Pro</h1>
</div>
<div class="sub-title">Sistema Inteligente de Gestão de ASO</div>
""", unsafe_allow_html=True)

# =========================================================
# LÓGICA DE DADOS (MAPEAMENTO RESILIENTE)
# =========================================================

SHEET_ID = "1G_oVT9gK-n_jGh5R4g65qUwK_MfQGvCX-SA4NHNNflU"

UNIDADES = {
    "C-CHP": "1201656811",  # Chapecó
    "C-CTBA": "145843404",   # Curitiba
    "S-SJ": "1712391604",    # São José
    "D-MG": "1071212860",    # Minas Gerais
    "D-MS": "1260438020",    # Mato Grosso do Sul
    "D-SP": "1323067532"     # São Paulo
}

@st.cache_data(ttl=15)
def load_data(gid):
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=csv&gid={str(gid).strip()}"
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as resp:
            df = pd.read_csv(resp)

        # 1. Limpeza de espaços em branco e quebras de linha em todos os cabeçalhos
        df.columns = [str(c).replace("\n", " ").strip() for c in df.columns]

        # 2. Padronização automática de cabeçalhos
        for c in list(df.columns):
            c_upper = c.upper()
            if c_upper in ["VENC", "VENCIMENTO", "VENC.", "DATA VENCIMENTO", "DATA DE VENCIMENTO"]:
                df.rename(columns={c: "Venc"}, inplace=True)
            elif c_upper in ["NOME", "COLABORADOR", "FUNCIONARIO", "NOME COMPLETO"]:
                df.rename(columns={c: "Nome"}, inplace=True)
            elif c_upper in ["DATA ASO", "ASO", "ULTIMO ASO", "DATA DO ASO"]:
                df.rename(columns={c: "Data ASO"}, inplace=True)
            elif c_upper in ["CARGO", "FUNCAO"]:
                df.rename(columns={c: "Cargo"}, inplace=True)
            elif c_upper in ["SETOR", "DEPARTAMENTO", "DPTO"]:
                df.rename(columns={c: "Setor"}, inplace=True)
            elif c_upper in ["AGENDAMENTO", "STATUS AGENDAMENTO"]:
                df.rename(columns={c: "Agendamento"}, inplace=True)

        if "Nome" in df.columns:
            df = df.dropna(subset=["Nome"]).copy()
            df["Nome"] = df["Nome"].astype(str).str.strip()
            df = df[df["Nome"] != ""]
            df = df[df["Nome"].str.lower() != "nan"]
            df = df[df["Nome"].str.lower() != "none"]

        return df
    except Exception:
        return pd.DataFrame()

@st.cache_data(ttl=15)
def carregar_base_completa():
    frames = []
    for u_nome, gid in UNIDADES.items():
        d = load_data(gid)
        if not d.empty and "Nome" in d.columns:
            d["UNIDADE_REAL"] = u_nome
            frames.append(d)
    if frames:
        return pd.concat(frames, ignore_index=True)
    return pd.DataFrame()

def gerar_docx(dados, tipo, data_sugestao, unidade_padrao="MACROMAQ"):
    doc = Document()
    if PATH_LOGO_DOC and os.path.exists(PATH_LOGO_DOC):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(PATH_LOGO_DOC, width=Inches(1.2))

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run(f"FORMULÁRIO PARA AGENDAMENTO {tipo}")
    run.bold, run.font.size = True, Pt(16)

    table = doc.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    labels = ["Nome Completo", "Cargo", "Setor", "Unidade", "Cliente", "Local", "Data Sugestão"]

    nome_val = str(dados.get("Nome", "")).strip() if pd.notna(dados.get("Nome")) else ""
    cargo_val = str(dados.get("Cargo", "")).strip() if pd.notna(dados.get("Cargo")) else ""
    setor_val = str(dados.get("Setor", "")).strip() if pd.notna(dados.get("Setor")) else ""
    unid_val = str(dados.get("UNIDADE", dados.get("UNIDADE_REAL", unidade_padrao))).strip()

    valores = [
        nome_val,
        cargo_val,
        setor_val,
        unid_val,
        "MACROMAQ",
        "Arapoti",
        data_sugestao.strftime("%d/%m/%Y")
    ]

    for i, (l, v) in enumerate(zip(labels, valores)):
        table.rows[i].cells[0].text, table.rows[i].cells[1].text = l, v
        table.rows[i].cells[0].paragraphs[0].runs[0].bold = True

    target = BytesIO()
    doc.save(target)
    return target.getvalue()

def renderizar_card_e_agendamento(row, unidade_nome, idx_chave, banco_ag, hoje):
    nome_str = str(row["Nome"]).strip()
    cargo_str = str(row.get("Cargo", "Não informado")) if pd.notna(row.get("Cargo")) and str(row.get("Cargo")).strip() != "" else "Não informado"
    setor_str = str(row.get("Setor", "Não informado")) if pd.notna(row.get("Setor")) and str(row.get("Setor")).strip() != "" else "Não informado"

    data_aso_val = str(row.get("Data ASO", "")).strip()

    venc_dt = pd.to_datetime(row.get("Venc"), dayfirst=True, errors="coerce")
    if pd.notna(venc_dt):
        dias = int((venc_dt - hoje).days)
        venc_str = venc_dt.strftime("%d/%m/%Y")
    else:
        dias = 9999
        venc_str = "Não informada"

    chave_registro = f"{unidade_nome}_{nome_str}_{venc_str.replace('/', '-')}"
    dados_locais = banco_ag.get(chave_registro, None)

    agendado_ativo = False
    texto_detalhe_agendamento = ""

    if dados_locais is not None:
        agendado_ativo = bool(dados_locais.get("status", False))
        data_salva = dados_locais.get("data", "")
        hora_salva = dados_locais.get("hora", "")
        obs_salva = dados_locais.get("obs", "")
        if agendado_ativo:
            texto_detalhe_agendamento = f"🗓 <b>Agendado:</b> {data_salva}" + (f" às {hora_salva}" if hora_salva else "")
            if obs_salva:
                texto_detalhe_agendamento += f"<br>📝 <b>Obs:</b> {obs_salva}"
    else:
        if "Agendamento" in row.index and pd.notna(
