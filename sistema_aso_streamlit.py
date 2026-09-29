import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
from datetime import datetime, timedelta
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
import os
import json
import urllib.request
from PIL import Image
import base64
import unicodedata
from urllib.parse import unquote

# =========================================================
# CONFIGURAÇÃO DE CAMINHOS E WEBHOOK DO APPS SCRIPT
# =========================================================

BASE_DIR = os.path.dirname(__file__) if "__file__" in locals() else "."

# Cole aqui o link da sua Web App do Google Apps Script
APPS_SCRIPT_WEBHOOK_URL = "" 

def localizar_arquivo(caminho_local, nome_arquivo):
    if os.path.exists(caminho_local):
        return caminho_local
    caminho_projeto = os.path.join(BASE_DIR, nome_arquivo)
    return caminho_projeto if os.path.exists(caminho_projeto) else None

PATH_ASO_IMG = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\ASO.png", "ASO.png")
PATH_LOGO = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\logo.png", "logo.png")
PATH_LOGO_DOC = localizar_arquivo(r"C:\Users\dilceu.gomes\Desktop\sistema_aso\adivitta.png", "adivitta.png")
PATH_AGENDAMENTOS_JSON = os.path.join(BASE_DIR, "agendamentos_persistentes.json")

def carregar_imagem_base64(path):
    if path and os.path.exists(path):
        with open(path, "rb") as img_file:
            return base64.b64encode(img_file.read()).decode()
    return None

def normalizar_texto(texto):
    if not isinstance(texto, str):
        return ""
    texto = unicodedata.normalize("NFKD", texto).encode("ASCII", "ignore").decode("ASCII")
    return texto.strip().upper()

# =========================================================
# PERSISTÊNCIA LOCAL E ENVIO PARA O GOOGLE APPS SCRIPT
# =========================================================

def ler_banco_agendamentos():
    if os.path.exists(PATH_AGENDAMENTOS_JSON):
        try:
            with open(PATH_AGENDAMENTOS_JSON, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def gravar_agendamento(chave, dados):
    banco = ler_banco_agendamentos()
    banco[chave] = dados
    try:
        with open(PATH_AGENDAMENTOS_JSON, "w", encoding="utf-8") as f:
            json.dump(banco, f, ensure_ascii=False, indent=2)
    except Exception as e:
        st.error(f"Erro ao salvar localmente: {e}")

def enviar_webhook_planilha(payload):
    """Envia alteração para o script doPost da planilha Google."""
    if not APPS_SCRIPT_WEBHOOK_URL or not APPS_SCRIPT_WEBHOOK_URL.startswith("http"):
        return
    try:
        dados_json = json.dumps(payload).encode("utf-8")
        req = urllib.request.Request(APPS_SCRIPT_WEBHOOK_URL, data=dados_json, headers={"Content-Type": "application/json"})
        urllib.request.urlopen(req, timeout=8)
    except Exception:
        pass

def remover_agendamento(chave, payload_sync=None):
    banco = ler_banco_agendamentos()
    banco[chave] = {"status": False, "data": "", "hora": "", "obs": ""}
    try:
        with open(PATH_AGENDAMENTOS_JSON, "w", encoding="utf-8") as f:
            json.dump(banco, f, ensure_ascii=False, indent=2)
    except Exception as e:
        st.error(f"Erro ao atualizar: {e}")
    if payload_sync:
        enviar_webhook_planilha(payload_sync)

# =========================================================
# CONFIGURAÇÃO DA PÁGINA
# =========================================================

try:
    if PATH_ASO_IMG:
        page_icon = Image.open(PATH_ASO_IMG)
    else:
        page_icon = "🛡️"
except Exception:
    page_icon = "🛡️"

st.set_page_config(
    page_title="Gestão ASO Pro",
    layout="wide",
    page_icon=page_icon
)

# =========================================================
# CSS GLOBAL E TOPO
# =========================================================

aso_base64 = carregar_imagem_base64(PATH_ASO_IMG)

st.markdown(f"""
<style>
[data-testid="stSidebarNav"] {{display: none !important;}}
.stApp {{ background-color: #f1f5f9; }}
section[data-testid="stSidebar"] {{ background: linear-gradient(180deg, #8390a8, #1e293b); }}
section[data-testid="stSidebar"] label, section[data-testid="stSidebar"] p {{ color: white !important; }}
div[data-testid="metric-container"] {{ background: white; border-radius: 18px; padding: 15px; border: 1px solid #e2e8f0; box-shadow: 0 4px 18px rgba(0,0,0,0.06); }}

.stDownloadButton button {{ width: 100%; background: linear-gradient(90deg, #2563eb, #1d4ed8); color: white; border-radius: 12px; font-weight: bold; }}

section[data-testid="stSidebar"] .stButton > button {{
    background-color: #2563eb !important;
    color: white !important;
    border-radius: 12px !important;
    border: 1px solid rgba(255,255,255,0.3) !important;
    font-weight: bold !important;
    height: 45px !important;
    transition: 0.3s;
}}

.footer {{ position: fixed; left: 0; bottom: 0; width: 100%; background: rgba(0,0,0,0.8); color: white; text-align: center; padding: 10px; font-size: 13px; z-index: 999; }}

.menu-titulo {{
    color: white;
    font-weight: bold;
    font-size: 14px;
    margin-top: 20px;
    margin-bottom: 10px;
    border-bottom: 1px solid rgba(255,255,255,0.2);
    padding-bottom: 5px;
}}

.header-wrapper {{
    display: flex;
    align-items: center;
    gap: 20px;
    margin-bottom: 5px;
}}
.main-title {{
    font-size: 40px;
    font-weight: 800;
    color: #111827;
    margin: 0;
}}
.sub-title {{
    color: #64748b;
    margin-bottom: 25px;
    margin-left: 120px;
}}
</style>

<div class="header-wrapper">
    {"<img src='data:image/png;base64," + aso_base64 + "' width='100'>" if aso_base64 else "🛡️"}
    <h1 class="main-title">Gestão ASO Pro</h1>
</div>
<div class="sub-title">Sistema Inteligente de Gestão de ASO</div>
""", unsafe_allow_html=True)

# =========================================================
# LÓGICA DE DADOS
# =========================================================

SHEET_ID = "1G_oVT9gK-n_jGh5R4g65qUwK_MfQGvCX-SA4NHNNflU"

UNIDADES = {
    "S-SJ": "1712391604",
    "J-CTBA": "145843404",
    "D-MG": "1071212860",
    "D-ITU": "1323067532",
    "J-CHP": "1549718037",
    "D-SP": "0"
}

@st.cache_data(ttl=30)
def load_data(gid):
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=csv&gid={str(gid).strip()}"
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as resp:
            df = pd.read_csv(resp)
            
        df.columns = df.columns.astype(str).str.strip()
        
        if "Nome" in df.columns:
            df = df.dropna(subset=["Nome"]).copy()
            df["Nome"] = df["Nome"].astype(str).str.strip()
            df = df[df["Nome"] != ""]
            df = df[df["Nome"].str.lower() != "nan"]
            df = df[df["Nome"].str.lower() != "none"]
            
        return df
    except Exception:
        return pd.DataFrame()

@st.cache_data(ttl=30)
def carregar_base_completa():
    frames = []
    for u_nome, gid in UNIDADES.items():
        d = load_data(gid)
        if not d.empty and "Nome" in d.columns:
            d["UNIDADE_REAL"] = u_nome
            frames.append(d)
    if frames:
        return pd.concat(frames, ignore_index=True)
    return pd.DataFrame()

def gerar_docx(dados, tipo, data_sugestao, unidade_padrao="MACROMAQ"):
    doc = Document()
    if PATH_LOGO_DOC and os.path.exists(PATH_LOGO_DOC):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(PATH_LOGO_DOC, width=Inches(1.2))

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run(f"FORMULÁRIO PARA AGENDAMENTO {tipo}")
    run.bold, run.font.size = True, Pt(16)

    table = doc.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    labels = ["Nome Completo", "Cargo", "Setor", "Unidade", "Cliente", "Local", "Data Sugestão"]
    
    nome_val = str(dados.get("Nome", "")).strip() if pd.notna(dados.get("Nome")) else ""
    cargo_val = str(dados.get("Cargo", "")).strip() if pd.notna(dados.get("Cargo")) else ""
    setor_val = str(dados.get("Setor", "")).strip() if pd.notna(dados.get("Setor")) else ""
    unid_val = str(dados.get("UNIDADE", dados.get("UNIDADE_REAL", unidade_padrao))).strip()

    valores = [
        nome_val,
        cargo_val,
        setor_val,
        unid_val,
        "MACROMAQ",
        "Arapoti",
        data_sugestao.strftime("%d/%m/%Y")
    ]

    for i, (l, v) in enumerate(zip(labels, valores)):
        table.rows[i].cells[0].text, table.rows[i].cells[1].text = l, v
        table.rows[i].cells[0].paragraphs[0].runs[0].bold = True

    target = BytesIO()
    doc.save(target)
    return target.getvalue()

def renderizar_card_e_agendamento(row, unidade_nome, idx_chave, banco_ag, hoje):
    """Renderiza o ASO independente e permite alterar a Data ASO da Coluna H."""
    nome_str = str(row["Nome"]).strip()
    cargo_str = str(row.get("Cargo", "Não informado")) if pd.notna(row.get("Cargo")) and str(row.get("Cargo")).strip() != "" else "Não informado"
    setor_str = str(row.get("Setor", "Não informado")) if pd.notna(row.get("Setor")) and str(row.get("Setor")).strip() != "" else "Não informado"
    
    # Tratamento da Data do ASO Atual (Coluna H)
    data_aso_val = str(row.get("Data ASO", "")).strip()
    
    # Tratamento da Data de Vencimento (Coluna I)
    venc_dt = pd.to_datetime(row.get("Venc"), dayfirst=True, errors="coerce")
    if pd.notna(venc_dt):
        dias = int((venc_dt - hoje).days)
        venc_str = venc_dt.strftime("%d/%m/%Y")
    else:
        dias = 9999
        venc_str = "Não informada"

    # Chave única por ASO (Nome + Vencimento)
    chave_registro = f"{unidade_nome}_{nome_str}_{venc_str.replace('/', '-')}"
    dados_locais = banco_ag.get(chave_registro, None)

    agendado_ativo = False
    texto_detalhe_agendamento = ""
    
    if dados_locais is not None:
        agendado_ativo = bool(dados_locais.get("status", False))
        data_salva = dados_locais.get("data", "")
        hora_salva = dados_locais.get("hora", "")
        obs_salva = dados_locais.get("obs", "")
        if agendado_ativo:
            texto_detalhe_agendamento = f"🗓 <b>Agendado:</b> {data_salva}" + (f" às {hora_salva}" if hora_salva else "")
            if obs_salva:
                texto_detalhe_agendamento += f"<br>📝 <b>Obs:</b> {obs_salva}"
    else:
        if "Agendamento" in row.index and pd.notna(row["Agendamento"]):
            txt_planilha = str(row["Agendamento"]).strip()
            if txt_planilha and txt_planilha.lower() not in ["nan", "none"]:
                agendado_ativo = True
                texto_detalhe_agendamento = f"🗓️ <b>Agendado:</b> {txt_planilha}"

    cor, fundo, status = (
        ("#ef4444", "#fff1f2", "🚨 ASO VENCIDO") if dias < 0 
        else (("#f59e0b", "#fff7ed", f"⚠️ Vence em {dias} dias") if dias <= 3 
        else ("#10b981", "#ecfdf5", f"✅ Vence em {dias} dias"))
    )

    badge_html = ""
    detalhes_html = ""
    if agendado_ativo:
        badge_html = '<span style="background:#2563eb; color:white; padding:4px 10px; border-radius:8px; font-size:12px; font-weight:bold; margin-left:10px;">📌 AGENDADO</span>'
        detalhes_html = f"""
        <div style="margin-top:10px; padding:10px 14px; background:#e0f2fe; border-radius:10px; color:#0369a1; font-size:13px; border:1px solid #bae6fd;">
            {texto_detalhe_agendamento}
        </div>
        """

    html_card = f"""
    <div style="background:{fundo}; border-left:8px solid {cor}; border-radius:18px; padding:22px; margin-bottom:15px; box-shadow:0 4px 18px rgba(0,0,0,0.08); font-family:Arial;">
        <div style="font-size:22px; font-weight:700; color:#111827; margin-bottom:10px; display:flex; align-items:center; flex-wrap:wrap;">
            {nome_str} {badge_html}
        </div>
        <div style="color:#475569; font-size:15px; line-height:1.8;">
            👔 <b>Cargo:</b> {cargo_str}<br>
            🏭 <b>Setor:</b> {setor_str}<br>
            🩺 <b>Data ASO (Col. H):</b> {data_aso_val if data_aso_val else 'Não informada'}<br>
            📅 <b>Vencimento (Col. I):</b> {venc_str}
        </div>
        {detalhes_html}
        <div style="margin-top:15px; font-size:16px; font-weight:bold; color:{cor};">{status}</div>
    </div>"""

    altura_card = 330 if agendado_ativo else 260
    components.html(html_card, height=altura_card)

    partes_nome = nome_str.split()
    primeiro_nome_card = partes_nome[0] if len(partes_nome) > 0 else "Colaborador"

    tab_agendamento, tab_atualizar_h = st.tabs([f"⚙️ Agendar / Emitir Guia", f"🔄 Atualizar Data do ASO (Col. H)"])

    with tab_agendamento:
        tipo = st.selectbox("Tipo de Exame", ["PERIÓDICO", "MUDANÇA DE RISCO", "RETORNO"], key=f"t_{idx_chave}")

        padrao_data = (hoje + timedelta(days=2)).date()
        if dados_locais and dados_locais.get("data"):
            try:
                padrao_data = datetime.strptime(dados_locais["data"], "%d/%m/%Y").date()
            except Exception:
                pass

        dt_sugestao = st.date_input("📅 Data Sugerida / Agendamento", value=padrao_data, key=f"d_{idx_chave}")
        padrao_hora = dados_locais.get("hora", "08:00") if dados_locais else "08:00"
        hora_input = st.text_input("⏰ Horário (ex: 08:30)", value=padrao_hora, key=f"hr_{idx_chave}")

        padrao_obs = dados_locais.get("obs", "") if dados_locais else (str(row.get("Agendamento", "")) if pd.notna(row.get("Agendamento")) and str(row.get("Agendamento")).lower() not in ["nan", "none"] else "")
        obs_input = st.text_area("📝 Observações (Clínica, médico, detalhes...)", value=padrao_obs, key=f"obs_{idx_chave}")

        def salvar_e_agendar(chave=chave_registro, data_val=dt_sugestao, hora_val=hora_input, obs_val=obs_input):
            dados_salvar = {
                "status": True,
                "data": data_val.strftime("%d/%m/%Y"),
                "hora": hora_val.strip(),
                "obs": obs_val.strip(),
                "unidade": unidade_nome,
                "venc_original": venc_str
            }
            gravar_agendamento(chave, dados_salvar)
            
            txt_status_planilha = f"{data_val.strftime('%d/%m/%Y')} às {hora_val.strip()}" + (f" - {obs_val.strip()}" if obs_val.strip() else "")
            enviar_webhook_planilha({
                "aba": unidade_nome,
                "nome": nome_str,
                "venc_original": venc_str,
                "status": txt_status_planilha,
                "acao": "agendamento"
            })

        btn_doc = gerar_docx(row, tipo, dt_sugestao, unidade_nome)

        st.download_button(
            label="📥 Baixar Documento e Marcar como Agendado",
            data=btn_doc,
            file_name=f"ASO_{nome_str}_{venc_str.replace('/', '-')}.docx",
            key=f"b_{idx_chave}",
            on_click=salvar_e_agendar
        )

        c_salvar, c_limpar = st.columns(2)
        with c_salvar:
            if st.button("💾 Apenas Salvar Status", key=f"btn_salvar_{idx_chave}"):
                salvar_e_agendar()
                st.toast("✅ Agendamento salvo!")
                st.rerun()

        with c_limpar:
            if agendado_ativo:
                if st.button("🗑️ Limpar / Desmarcar", key=f"btn_limpar_{idx_chave}"):
                    remover_agendamento(chave_registro, {
                        "aba": unidade_nome,
                        "nome": nome_str,
                        "venc_original": venc_str,
                        "acao": "limpar"
                    })
                    st.toast("🗑️ Agendamento removido!")
                    st.rerun()

    # =========================================================
    # ABA NOVA: ALTERAR A DATA DA COLUNA H (DATA ASO) NA PLANILHA
    # =========================================================
    with tab_atualizar_h:
        st.markdown(f"**Alterar 'Data ASO' de {primeiro_nome_card}**")
        st.caption("Esta nova data será gravada diretamente na **Coluna H**, e a planilha recalculará o vencimento automaticamente na Coluna I.")

        nova_data_coluna_h = st.date_input("Nova Data do ASO (Coluna H)", value=hoje.date(), key=f"dt_col_h_{idx_chave}")

        if st.button("💾 Gravar Nova Data na Coluna H", key=f"btn_gravar_h_{idx_chave}"):
            # Remove o status de agendado localmente (pois o ASO foi feito)
            remover_agendamento(chave_registro)
            
            # Envia para a planilha atualizar a coluna H e limpar a coluna M
            enviar_webhook_planilha({
                "aba": unidade_nome,
                "nome": nome_str,
                "venc_original": venc_str,
                "nova_data_aso": nova_data_coluna_h.strftime("%d/%m/%Y"),
                "acao": "atualizar_data_aso"
            })
            st.cache_data.clear()
            st.success(f"✅ Data {nova_data_coluna_h.strftime('%d/%m/%Y')} gravada na Coluna H com sucesso!")
            st.rerun()

# =========================================================
# SIDEBAR
# =========================================================

if PATH_LOGO and os.path.exists(PATH_LOGO):
    st.sidebar.image(PATH_LOGO, width=300)

aba_nome = st.sidebar.selectbox(
    "🏢 Selecione a unidade",
    list(UNIDADES.keys())
)

st.sidebar.success("✅ Sistema Online")
st.sidebar.markdown("<br>", unsafe_allow_html=True)
st.sidebar.markdown("<div class='menu-titulo'>⚡ ACESSO RÁPIDO</div>", unsafe_allow_html=True)

if st.sidebar.button("📋 CHECKLIST SSMA"):
    st.switch_page("pages/app_ssma_ia.py")

# =========================================================
# 🔍 BUSCA ATIVA E EMISSÃO DE GUIA AVULSA
# =========================================================
st.markdown("### 🔍 Busca Ativa e Emissão de Guia Avulsa")

df_todos = carregar_base_completa()
banco_ag = ler_banco_agendamentos()
hoje = datetime.now()

if not df_todos.empty and "Nome" in df_todos.columns:
    colab_param = st.query_params.get("colaborador", None)
    
    lista_todos_nomes = sorted([
        str(n).strip() for n in df_todos["Nome"].dropna().unique() 
        if str(n).strip() and str(n).lower() not in ["nan", "none"]
    ])
    
    nome_pre_selecionado = None
    if colab_param:
        nome_param_limpo = normalizar_texto(unquote(str(colab_param)))
        for n in lista_todos_nomes:
            if normalizar_texto(n) == nome_param_limpo or nome_param_limpo in normalizar_texto(n):
                nome_pre_selecionado = n
                break
        
        if not nome_pre_selecionado:
            nome_pre_selecionado = unquote(str(colab_param)).strip()

    if nome_pre_selecionado:
        if nome_pre_selecionado in lista_todos_nomes:
            opcoes_finais = [nome_pre_selecionado] + [x for x in lista_todos_nomes if x != nome_pre_selecionado]
        else:
            opcoes_finais = [nome_pre_selecionado] + lista_todos_nomes
        idx_padrao = 0
    else:
        opcoes_finais = ["Selecione um colaborador..."] + lista_todos_nomes
        idx_padrao = 0

    colab_escolhido = st.selectbox(
        "Digite ou selecione o nome do colaborador:",
        options=opcoes_finais,
        index=idx_padrao
    )

    if colab_escolhido and colab_escolhido != "Selecione um colaborador...":
        registros_colab = df_todos[df_todos["Nome"].astype(str) == str(colab_escolhido)]
        if not registros_colab.empty:
            for b_idx, (_, d_row) in enumerate(registros_colab.iterrows()):
                unid_row = str(d_row.get("UNIDADE_REAL", aba_nome))
                v_str = str(d_row.get("Venc", ""))
                st.markdown(f"##### 👤 Ficha e Agendamento: {colab_escolhido} (Venc: {v_str})")
                renderizar_card_e_agendamento(d_row, unid_row, f"busca_ativa_{b_idx}", banco_ag, hoje)
        else:
            mock_row = pd.Series({
                "Nome": str(colab_escolhido),
                "Cargo": "Geral",
                "Setor": "Operacional",
                "UNIDADE_REAL": aba_nome,
                "Venc": hoje + timedelta(days=30)
            })
            renderizar_card_e_agendamento(mock_row, aba_nome, "busca_ativa_mock", banco_ag, hoje)

st.markdown("<hr style='margin: 25px 0;'>", unsafe_allow_html=True)

# =========================================================
# MONITORAMENTO DE ALERTAS COM FILTRO DE VENCIMENTOS
# =========================================================
try:
    df_unidade = load_data(UNIDADES[aba_nome])
    
    if not df_unidade.empty and "Nome" in df_unidade.columns and "Venc" in df_unidade.columns:
        df_unidade["Venc"] = pd.to_datetime(df_unidade["Venc"], dayfirst=True, errors="coerce")
        df_base = df_unidade.dropna(subset=["Venc", "Nome"]).copy()
        
        # Filtros de higienização
        df_base["Nome"] = df_base["Nome"].astype(str).str.strip()
        df_base = df_base[df_base["Nome"] != ""]
        df_base = df_base[df_base["Nome"].str.lower() != "nan"]
        df_base = df_base[df_base["Nome"].str.lower() != "none"]
        df_base = df_base[df_base["Venc"].dt.year >= 2020]
        
        df_base["Dias"] = (df_base["Venc"] - hoje).dt.days

        # Identifica agendamentos respeitando Nome + Vencimento
        def esta_agendado_check(row):
            v_s = row["Venc"].strftime("%d-%m-%Y")
            chave = f"{aba_nome}_{row['Nome']}_{v_s}"
            d = banco_ag.get(chave)
            if d and d.get("status"):
                return True
            return False

        df_base["Is_Agendado"] = df_base.apply(esta_agendado_check, axis=1)

        total_vencidos = len(df_base[df_base["Dias"] < 0])
        total_agendados = len(df_base[df_base["Is_Agendado"]])

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("🏢 Unidade", aba_nome)
        c2.metric("👥 Colaboradores", len(df_unidade))
        c3.metric("🚨 Vencidos", total_vencidos)
        c4.metric("📌 Agendados", total_agendados)

        st.markdown("#### 📋 Monitoramento e Prazos")

        filtro_selecionado = st.radio(
            "Filtrar colaboradores:",
            options=[
                "🚨 Vencidos e Próximos (Padrão: até 30 dias)",
                "🚨 Apenas Vencidos",
                "⚠️ Vence em até 10 dias",
                "📅 Vence em até 60 dias",
                "🌐 Todos a Vencer (+30 dias)",
                "📌 Somente Agendados",
                "📄 Todos os Colaboradores"
            ],
            horizontal=True
        )

        if filtro_selecionado == "🚨 Vencidos e Próximos (Padrão: até 30 dias)":
            alertas = df_base[(df_base["Venc"] <= hoje + timedelta(days=30)) | (df_base["Is_Agendado"])].copy()
        elif filtro_selecionado == "🚨 Apenas Vencidos":
            alertas = df_base[df_base["Dias"] < 0].copy()
        elif filtro_selecionado == "⚠️ Vence em até 10 dias":
            alertas = df_base[(df_base["Dias"] <= 10) | (df_base["Is_Agendado"])].copy()
        elif filtro_selecionado == "📅 Vence em até 60 dias":
            alertas = df_base[(df_base["Dias"] <= 60) | (df_base["Is_Agendado"])].copy()
        elif filtro_selecionado == "🌐 Todos a Vencer (+30 dias)":
            alertas = df_base[df_base["Dias"] > 30].copy()
        elif filtro_selecionado == "📌 Somente Agendados":
            alertas = df_base[df_base["Is_Agendado"]].copy()
        else:
            alertas = df_base.copy()

        alertas = alertas.sort_values(by="Venc")

        if alertas.empty:
            st.info("Nenhum colaborador encontrado para o filtro selecionado.")
        else:
            st.caption(f"Mostrando **{len(alertas)}** registos de ASO.")
            cols = st.columns(2)
            for idx, (_, row) in enumerate(alertas.iterrows()):
                with cols[idx % 2]:
                    renderizar_card_e_agendamento(row, aba_nome, f"painel_{idx}", banco_ag, hoje)
    else:
        st.warning(f"Sem dados carregados para a unidade {aba_nome}.")
except Exception as e:
    st.error(f"Erro ao processar dados: {e}")

st.markdown("""<div class="footer">© 2026 Gestão Documentos | Desenvolvido por: Dilceu Junior</div>""", unsafe_allow_html=True)
